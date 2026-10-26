import contextlib
import copy
import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from docx import Document

import upload_language as app
from modules import extraction, languages, text_processing


def game_template(name="game-one"):
    return {
        "fontName": name + " font", "logo": {"path": name + "_logo", "offset": "0,300"},
        "background": {"path": name + "_background"}, "payTables": [{"images": [{"src": "sym_8a.png"}]}],
        "payLines": [{"titles": [{"text": "50 LINES"}]}], "customGameData": {"game": name},
        "rules": {
            "settings": {"items": {"font": {"size": 19, "color": "white"}, "offset": "0,170", "gap": 28},
                         "titles": {"font": {"size": 32}}, "symbols": {"default": {"scale": 0.4}}},
            "pages": [
                {"titles": [{"text": "GAME RULES", "color": "yellow"}],
                 "items": {"setting": {"font": {"size": 21}}, "lines": [
                     "• First ##sym_8a.png## pays.", "", "• Keep this missing rule.", "", "• Final rule."]}},
                {"images": [{"path": "reference_picture", "offset": "0,0"}]},
                {"titles": [{"text": "NEW FEATURE"}], "items": {"lines": ["• Pick a chest."]}},
            ]},
    }


def write_docx(path, sections, english_header="INGLÉS", continuation=False):
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for source_title, target_title, entries in sections:
        rows = [("", english_header, "GERMAN"), ("", source_title, target_title)]
        rows.extend((str(i), source, target) for i, (source, target) in enumerate(entries, 1))
        tables = [rows[:-1], [rows[-1]]] if continuation and len(entries) > 1 else [rows]
        for values in tables:
            table = document.add_table(rows=len(values), cols=3)
            for row, texts in zip(table.rows, values):
                for cell, text in zip(row.cells, texts):
                    cell.text = text
    document.save(path)


class ImporterTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.project = Path(self.temp.name)
        self.root = self.project / "uploadLanguage"
        self.templates = self.root / "template"
        self.documents = self.root / "files"
        self.documents.mkdir(parents=True)
        self.config = dict(documents="files", templates="template", output="output",
                           contentWidth=650, symbolWidth=50, bottomY=-150, prefixes=languages.PREFIXES, tokens={})
        app.write_json(self.root / "config.json", self.config)
        self.english = game_template()
        self.jackpot = copy.deepcopy(self.english)
        self.jackpot["rules"]["pages"] = [{"titles": [{"text": "JACKPOT LINK"}],
            "items": {"lines": ["{fixedJackpot}• Each bet contributes."]}}]
        self.save_english()

    def save_english(self):
        app.write_json(self.templates / "info.json", self.english)
        app.write_json(self.templates / "jackpotpages.json", self.jackpot)

    def run_app(self, *args):
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            return app.main(["--config", str(self.root / "config.json"), *args])

    def add_german(self, entries=None, title="SPIELREGELN", filename="GERMANI.docx"):
        if entries is None:
            entries = [("• First [SYMBOL_8A] pays.", "• Erstes [SYMBOL_8A] gewinnt."),
                       ("• Final rule.", "• Letzte Regel.")]
        write_docx(self.documents / filename, [("GAME RULES", title, entries)])

    def generated(self, filename="info.json"):
        return app.read_json(self.root / "output/lang/deu/info" / filename)

    def report(self):
        return app.read_json(self.root / "output/report.json")

    def test_complete_prefix_mapping_and_arbitrary_name_suffix(self):
        for prefix, lang in languages.PREFIXES.items():
            (self.documents / f"{prefix.lower()}anything.DOCX").touch()
        found, ignored = extraction.find_documents(self.documents, self.config)
        self.assertEqual(set(found), set(languages.PREFIXES.values()))
        self.assertEqual(ignored, [])

    def test_duplicate_prefix_is_rejected(self):
        self.add_german(filename="GERMAN.docx")
        self.add_german(filename="GERMANI.docx")
        with self.assertRaisesRegex(ValueError, "Dos DOCX"):
            self.run_app()
        self.assertFalse((self.templates / "deu").exists())

    def test_document_prefix_ignores_leading_whitespace(self):
        self.add_german(filename=" GERMAN.docx")
        (self.documents / " ~$GERMAN.docx").write_text("temporary file", encoding="utf-8")
        found, ignored = extraction.find_documents(self.documents, self.config)
        self.assertEqual(found["deu"].name, " GERMAN.docx")
        self.assertEqual(ignored, [])
        self.assertEqual(self.run_app("--lang", "deu"), 0)

    def test_shared_documents_pdf_and_word_temporary_files_are_not_read(self):
        self.add_german()
        for name in ("ENG-ESP-CHIN.docx", "ESPAÑOL.docx", "CHINO.docx", "~$GERMAN.docx", "GERMAN.pdf"):
            (self.documents / name).write_text("not a DOCX", encoding="utf-8")
        self.assertEqual(self.run_app(), 0)
        manifest = app.read_json(self.root / "output/template-manifest.json")
        self.assertEqual(set(manifest["documents"]), {"deu"})
        self.assertIn("ENG-ESP-CHIN.docx", manifest["ignoredDocuments"])

    def test_google_drive_is_used_when_upload_has_no_documents(self):
        drive_url = "https://drive.google.com/drive/folders/example-folder"
        self.config["googleDriveUrl"] = drive_url
        app.write_json(self.root / "config.json", self.config)

        def download(*, id, output, **_):
            self.assertEqual(id, "german-id")
            path = Path(output)
            write_docx(path, [("GAME RULES", "SPIELREGELN", [
                ("• First [SYMBOL_8A] pays.", "• Erstes [SYMBOL_8A] gewinnt."),
                ("• Final rule.", "• Letzte Regel."),
            ])])
            return str(path)

        with mock.patch.object(extraction, "gdown") as gdown:
            gdown.download_folder.return_value = [SimpleNamespace(id="german-id", path="GERMANI.docx")]
            gdown.download.side_effect = download
            self.assertEqual(self.run_app(), 0)
        self.assertEqual(self.generated()["rules"]["pages"][0]["titles"][0]["text"], "SPIELREGELN")
        manifest = app.read_json(self.root / "output/template-manifest.json")
        self.assertEqual(manifest["documents"]["deu"], f"{drive_url} (GERMANI.docx)")

    def test_google_drive_download_failure_does_not_block_another_language(self):
        self.config["googleDriveUrl"] = "https://drive.google.com/drive/folders/example-folder"
        app.write_json(self.root / "config.json", self.config)

        def download(*, id, output, **_):
            if id == "german-id":
                raise RuntimeError("German download blocked")
            path = Path(output)
            write_docx(path, [("GAME RULES", "REGOLE", [
                ("• First [SYMBOL_8A] pays.", "• Prima [SYMBOL_8A] paga."),
                ("• Final rule.", "• Regola finale."),
            ])])
            return str(path)

        with mock.patch.object(extraction, "gdown") as gdown:
            gdown.download_folder.return_value = [
                SimpleNamespace(id="german-id", path="GERMAN.docx"),
                SimpleNamespace(id="italian-id", path="ITALIAN.docx"),
            ]
            gdown.download.side_effect = download
            self.assertEqual(self.run_app(), 1)
        self.assertTrue((self.root / "output/lang/ita/info/info.json").is_file())
        self.assertIn("German download blocked", self.report()[0]["message"])

    def test_upload_documents_take_priority_over_google_drive(self):
        self.add_german()
        self.config["googleDriveUrl"] = "https://drive.google.com/drive/folders/example-folder"
        app.write_json(self.root / "config.json", self.config)
        with mock.patch.object(extraction, "gdown") as gdown:
            self.assertEqual(self.run_app(), 0)
        gdown.download_folder.assert_not_called()

    def test_english_spanish_and_chinese_commands_are_excluded(self):
        for lang in languages.EXCLUDED:
            with self.subTest(lang=lang), self.assertRaisesRegex(ValueError, "excluidos"):
                self.run_app("--lang", lang)

    def test_complete_dynamic_bases_have_only_pages_emptied(self):
        self.add_german()
        self.assertEqual(self.run_app(), 0)
        for filename, original in (("info.json", self.english), ("jackpotpages.json", self.jackpot)):
            base = app.read_json(self.root / "output/templates" / filename)
            self.assertEqual(base, app.blank_template(original))
            self.assertEqual(base["rules"]["pages"], [])
            generated = self.generated(filename)
            self.assertEqual(app.blank_template(generated), base)
            self.assertEqual(app.read_json(self.templates / filename), original)

    def test_two_games_use_their_own_objects_and_sections(self):
        self.add_german()
        self.assertEqual(self.run_app(), 0)
        self.english = game_template("another-game")
        self.english["rules"]["pages"] = [{"titles": [{"text": "MYSTERY PICK"}],
            "items": {"lines": ["• Select a prize."]}}]
        self.save_english()
        write_docx(self.documents / "GERMANI.docx", [("MYSTERY PICK", "GEHEIME AUSWAHL", [("Select a prize.", "Wähle einen Preis.")])])
        self.assertEqual(self.run_app(), 0)
        generated = self.generated()
        self.assertEqual(app.blank_template(generated), app.blank_template(self.english))
        self.assertEqual(generated["rules"]["pages"][0]["titles"][0]["text"], "GEHEIME AUSWAHL")

    def test_missing_middle_rule_does_not_shift_later_translation(self):
        self.add_german()
        self.assertEqual(self.run_app(), 0)
        text = " ".join(line for page in self.generated()["rules"]["pages"] for line in page.get("items", {}).get("lines", []))
        self.assertLess(text.index("Erstes"), text.index("Keep this missing rule."))
        self.assertLess(text.index("Keep this missing rule."), text.index("Letzte Regel."))
        self.assertIn("##sym_8a.png##", text)
        self.assertTrue(any("regla 2" in entry["message"] for entry in self.report()))

    def test_position_matching_uses_translated_titles_and_rules(self):
        self.config["matching"] = "position"
        app.write_json(self.root / "config.json", self.config)
        self.english["rules"]["pages"] = [{"titles": [{"text": "HOLD & SPIN"}],
            "items": {"lines": ["• Hold and spin starts.", "• Hold and spin ends."]}}]
        self.save_english()
        write_docx(self.documents / "GERMAN.docx", [
            ("RESPIN FEATURE", "NEUER TITEL", [("Respin starts.", "Neue erste Regel."),
                                               ("Respin ends.", "Neue zweite Regel.")]),
            ("OTHER JACKPOT TITLE", "GEWINN", [("Different jackpot text.", "Neuer Gewinn.")])])
        self.assertEqual(self.run_app(), 0)
        pages = self.generated()["rules"]["pages"]
        self.assertEqual(pages[0]["titles"][0]["text"], "NEUER TITEL")
        self.assertIn("Neue erste Regel.", " ".join(pages[0]["items"]["lines"]))
        self.assertIn("Neue zweite Regel.", " ".join(pages[0]["items"]["lines"]))

    def test_position_matching_missing_row_falls_back_to_content(self):
        self.config["matching"] = "position"
        app.write_json(self.root / "config.json", self.config)
        self.test_missing_middle_rule_does_not_shift_later_translation()

    def test_missing_sections_and_visual_pages_are_preserved_exactly(self):
        self.add_german()
        self.assertEqual(self.run_app(), 0)
        generated = self.generated()
        self.assertIn(self.english["rules"]["pages"][1], generated["rules"]["pages"])
        self.assertIn(self.english["rules"]["pages"][2], generated["rules"]["pages"])
        self.assertEqual(self.generated("jackpotpages.json"), self.jackpot)

    def test_unknown_sections_do_not_create_new_content(self):
        write_docx(self.documents / "GERMAN.docx", [("OTHER GAME BONUS", "ANDERER BONUS", [("Other rule.", "Andere Regel.")])])
        self.assertEqual(self.run_app(), 0)
        self.assertEqual(self.generated(), self.english)
        self.assertTrue(any("no presente" in entry["message"] for entry in self.report()))

    def test_bad_translations_retain_english(self):
        for target in ("", "[SYMBOL_8A]", "• First [SYMBOL_8A] pays.", "[UNKNOWN] gewinnt.",
                       "[SYMBOL_18] gewinnt.", "[SYMBOL_9 gewinnt.", "• Eine. • Zwei."):
            with self.subTest(target=target):
                self.add_german([( "First [SYMBOL_8A] pays.", target)], title="GAME RULES")
                self.assertEqual(self.run_app(), 0)
                self.assertEqual(self.generated(), self.english)
                self.assertTrue(any("Inglés conservado" in entry["message"] for entry in self.report()))

    def test_marker_typos_are_repaired_when_the_marker_exists_in_the_english_row(self):
        for target in ("[SYMBOL_8A gewinnt.", "SYMBOL_8A] gewinnt.", "[[SYMBOL_8A]] gewinnt.",
                       "[SIMBOL_8A] gewinnt.", "[SÍMBOLO_8A] gewinnt."):
            with self.subTest(target=target):
                self.add_german([("First [SYMBOL_8A] pays.", target)], title="GAME RULES")
                self.assertEqual(self.run_app(), 0)
                self.assertEqual(self.generated()["rules"]["pages"][0]["items"]["lines"][0], "• ##sym_8a.png## gewinnt.")
                self.assertTrue(any("Marcador del DOCX corregido" in entry["message"] for entry in self.report()))

    def test_translated_jackpot_marker_name_is_resolved_from_the_english_row(self):
        self.english["rules"]["pages"] = [{"titles": [{"text": "GAME RULES"}],
            "items": {"lines": ["• The ##text_grand.png## awards a prize."]}}]
        self.save_english()
        self.add_german([("The [GRAND_JACKPOT_ICON] awards a prize.", "Das [GROßER _JACKPOT_ICON] gibt einen Preis.")])
        self.assertEqual(self.run_app(), 0)
        self.assertIn("Das ##text_grand.png## gibt", self.generated()["rules"]["pages"][0]["items"]["lines"][0])

    def test_text_inside_content_controls_and_merged_rows_is_read(self):
        """Google Docs exports put translations inside w:sdt and repeat vertically merged rows."""
        from docx.oxml import parse_xml
        path = self.documents / "GERMAN.docx"
        write_docx(path, [("GAME RULES", "SPIELREGELN", [("First [SYMBOL_8A] pays.", ""), ("Final rule.", "Letzte Regel.")])])
        doc = Document(path)
        cell = doc.tables[0].rows[2].cells[2]
        ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        cell._tc.remove(cell._tc.findall(extraction.W + "p")[0])
        cell._tc.append(parse_xml(f'<w:p {ns}><w:sdt><w:sdtContent><w:r><w:t>Erstes [SYMBOL_8A] gewinnt.</w:t>'
                                  '</w:r></w:sdtContent></w:sdt></w:p>'))
        table = doc.tables[0]
        new_row = copy.deepcopy(table.rows[2]._tr)
        table._tbl.insert(table._tbl.index(table.rows[2]._tr) + 1, new_row)
        for upper, lower in zip(table.rows[2].cells, table.rows[3].cells):
            upper._tc.get_or_add_tcPr().append(parse_xml(f'<w:vMerge {ns} w:val="restart"/>'))
            lower._tc.get_or_add_tcPr().append(parse_xml(f'<w:vMerge {ns}/>'))
        doc.save(path)
        rules = extraction.extract_docx(path, "deu", [])[text_processing.identity("GAME RULES")]["entries"]
        self.assertEqual([rule["text"] for rule in rules], ["Erstes [SYMBOL_8A] gewinnt.", "Letzte Regel."])

    def test_multiline_cell_is_split_into_template_sub_bullets(self):
        self.config["matching"] = "position"
        app.write_json(self.root / "config.json", self.config)
        self.english["rules"]["pages"] = [{"titles": [{"text": "GAME RULES"}], "items": {"lines": [
            "• Choose one:", "", "       • Collect.", "", "       • Multiply.", "", "• Final rule."]}}]
        self.save_english()
        self.add_german([("Choose one:\nCollect.\nMultiply.", "Wähle eins:\nSammeln.\nMultiplizieren."),
                         ("Final rule.", "Letzte Regel.")])
        self.assertEqual(self.run_app(), 0)
        self.assertEqual(self.generated()["rules"]["pages"][0]["items"]["lines"],
                         ["• Wähle eins:", "", "       • Sammeln.", "", "       • Multiplizieren.", "", "• Letzte Regel."])

    def test_docx_symbol_numbers_follow_the_template_assets(self):
        """The DOCX may number symbols differently (SYMBOL_15 = sym_16.png); learn it from aligned rows."""
        self.english["rules"]["pages"] = [{"titles": [{"text": "GAME RULES"}], "items": {"lines": [
            "• A ##sym_16.png## pays.", "", "• Extra rule.", "", "• Collect every ##sym_16.png##."]}}]
        self.save_english()
        self.add_german([("A [SYMBOL_15] pays.", "Ein [SYMBOL_15] zahlt."), ("Extra rule.", "Extra."),
                         ("Collect every [SYMBOL_15].", "Sammle jedes [SYMBOL_15].")])
        self.assertEqual(self.run_app(), 0)
        lines = self.generated()["rules"]["pages"][0]["items"]["lines"]
        self.assertIn("• Ein ##sym_16.png## zahlt.", lines)
        self.assertIn("• Sammle jedes ##sym_16.png##.", lines)

    def test_paytable_notes_translate_texts_outside_rules(self):
        self.english["payTables"] = [{"titles": [{"text": "KING"}, {"text": "CAN TRIGGER THE\nSPECIAL FEATURES"},
                                                 {"text": "OF A KIND"}]}]
        self.save_english()
        self.add_german()
        doc = Document(self.documents / "GERMANI.docx")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "WILD"
        table.rows[1].cells[0].text = "ENG: KING\nGERMAN: KÖNIG\n\nENG: CAN TRIGGER THE SPECIAL FEATURES\nGERMAN: KANN DIE SPEZIELLEN FUNKTIONEN AUSLÖSEN"
        body = doc.element.body
        body.insert(0, table._tbl)
        doc.save(self.documents / "GERMANI.docx")
        self.assertEqual(self.run_app(), 0)
        titles = [title["text"] for title in self.generated()["payTables"][0]["titles"]]
        self.assertEqual(titles[0], "KÖNIG")
        self.assertEqual(len(titles[1].split("\n")), 2)
        self.assertEqual(titles[1].replace("\n", " "), "KANN DIE SPEZIELLEN FUNKTIONEN AUSLÖSEN")
        self.assertEqual(titles[2], "OF A KIND")
        self.assertTrue(any("OF A KIND" in entry["message"] for entry in self.report()))

    def test_japanese_paytable_lines_keep_words_together(self):
        self.assertEqual(text_processing.balanced_lines("収集機能とフリーゲーム機能中に登場する", 3),
                         ["収集機能と", "フリーゲーム", "機能中に登場する"])
        self.assertEqual(text_processing.balanced_lines("ក\u200bខ\u200bគ\u200bឃ", 2), ["ក\u200bខ", "គ\u200bឃ"])

    def test_source_with_wrong_numeric_symbol_is_not_matched(self):
        self.add_german([("First [SYMBOL_18] pays.", "[SYMBOL_18] gewinnt.")], title="GAME RULES")
        self.assertEqual(self.run_app(), 0)
        self.assertEqual(self.generated(), self.english)

    def test_different_numeric_conditions_are_not_matched(self):
        for source, expected in (("Win 1.5 bets.", "Win 15 bets."), ("Win 3-4 bets.", "Win 34 bets."),
                                 ("Win -3 bets.", "Win 3 bets.")):
            with self.subTest(source=source):
                self.assertFalse(text_processing.source_match(source, expected, set(), {})[0])

    def test_ambiguous_duplicate_rules_retain_english(self):
        self.add_german([("Final rule.", "Eine Regel."), ("Final rule.", "Andere Regel.")], title="GAME RULES")
        self.assertEqual(self.run_app(), 0)
        self.assertEqual(self.generated(), self.english)

    def test_named_marker_can_be_inferred_from_matching_english_sentence(self):
        self.english["rules"]["pages"] = [{"titles": [{"text": "GAME RULES"}],
            "items": {"lines": ["• The ##text_grand.png## awards a prize."]}}]
        self.save_english()
        self.add_german([("The [GRAND_JACKPOT_ICON] awards a prize.", "Das [GRAND_JACKPOT_ICON] gibt einen Preis.")])
        self.assertEqual(self.run_app(), 0)
        text = " ".join(self.generated()["rules"]["pages"][0]["items"]["lines"])
        self.assertIn("Das ##text_grand.png##", text)

    def test_single_row_continuation_and_english_header_variants(self):
        for name in ("INGLÉS", "ENGLISH", "ENG"):
            with self.subTest(header=name):
                path = self.documents / "GERMAN.docx"
                write_docx(path, [("GAME RULES", "SPIELREGELN", [("First.", "Erste."), ("Final.", "Letzte.")])], name, True)
                extracted = extraction.extract_docx(path, "deu", [])
                self.assertEqual([entry["number"] for entry in extracted[text_processing.identity("GAME RULES")]["entries"]], [1, 2])

    def test_rule_numbers_are_assigned_by_order_including_continuations(self):
        for numbers in (("1", "2", "4", "5", "8"), ("1", "2", "", "4", "5"), ("",) * 5):
            with self.subTest(numbers=numbers):
                path = self.documents / "GERMAN.docx"
                entries = [(f"Rule {i}.", f"Regel {i}.") for i in range(1, 6)]
                write_docx(path, [("GAME RULES", "SPIELREGELN", entries)], continuation=True)
                doc = Document(path)
                rows = list(doc.tables[0].rows)[2:] + list(doc.tables[1].rows)
                for row, number in zip(rows, numbers):
                    row.cells[0].text = number
                doc.save(path)
                extracted = extraction.extract_docx(path, "deu", [])
                rules = extracted[text_processing.identity("GAME RULES")]["entries"]
                self.assertEqual([rule["number"] for rule in rules], [1, 2, 3, 4, 5])
                self.assertEqual([rule["text"] for rule in rules], [target for _, target in entries])

    def test_blank_target_header_and_nested_titles_are_imported(self):
        self.config["matching"] = "position"
        self.config["contentWidth"] = 1000
        self.config["bottomY"] = -500
        app.write_json(self.root / "config.json", self.config)
        self.english["rules"]["pages"] = [
            {"titles": [{"text": "NEW MAIN FEATURE"}, {"text": "FIRST SUBFEATURE", "color": "blue"}],
             "items": {"lines": ["• First rule.", "", "• Second rule."]}},
            {"titles": [{"text": "NEW MAIN FEATURE"}, {"text": "SECOND SUBFEATURE", "color": "green"}],
             "items": {"lines": ["• Third rule."]}}]
        self.save_english()
        path = self.documents / "GERMAN.docx"
        doc = Document()
        values = [("", "ENGLISH", ""), ("", "NEW MAIN FEATURE", "HAUPTFUNKTION"),
                  ("", "FIRST SUBFEATURE", "ERSTE FUNKTION"),
                  ("1", "First rule.", "Erste Regel."), ("2", "Second rule.", "Zweite Regel."),
                  ("", "SECOND SUBFEATURE", "ZWEITE FUNKTION"),
                  ("1", "Third rule.", "Dritte Regel.")]
        table = doc.add_table(rows=len(values), cols=3)
        for row, texts in zip(table.rows, values):
            for cell, text in zip(row.cells, texts):
                cell.text = text
        doc.save(path)
        self.assertEqual(self.run_app("--lang", "deu"), 0)
        pages = self.generated()["rules"]["pages"]
        self.assertEqual(len(pages), 2)
        self.assertEqual(pages[0]["titles"][1], {"text": "ERSTE FUNKTION", "color": "blue"})
        self.assertEqual(pages[1]["titles"][1], {"text": "ZWEITE FUNKTION", "color": "green"})
        self.assertEqual(pages[0]["items"]["lines"], ["• Erste Regel.", "", "• Zweite Regel."])
        self.assertEqual(pages[1]["items"]["lines"], ["• Dritte Regel."])
        extracted = app.read_json(self.root / "output/extracted.json")["deu"]
        section = extracted[text_processing.identity("NEW MAIN FEATURE")]
        self.assertEqual(len(section["entries"]), 3)
        self.assertEqual(len(section["subtitles"]), 2)

    def test_positional_symbols_use_english_assets_and_reject_target_changes(self):
        expected = "• The ##sym_8c.png## awards ##text_grand.png##."
        source = "The [SYMBOL_9c] awards [GRAND_JACKPOT_ICON]."
        mapping = text_processing.positional_token_mapping(source, expected)
        allowed = {"sym_8c.png", "text_grand.png"}
        self.assertEqual(text_processing.convert_tokens("[SYMBOL_9c] gewinnt [GRAND_ JACKPOT].", allowed, {}, mapping),
                         "##sym_8c.png## gewinnt ##text_grand.png##.")
        section = dict(title="FEATURE", rules=[expected])
        document = dict(title="FUNKTION", entries=[dict(number=1, english=source,
                                                        text="[SYMBOL_9c] gewinnt [SYMBOL_9c].")])
        report = []
        translations, _ = app.translate_rules(section, document, "deu", "test.docx",
                                              {"matching": "position"}, allowed, report)
        self.assertEqual(translations, [None])
        self.assertIn("símbolos no coinciden", report[0]["message"])
        self.assertEqual(text_processing.positional_token_mapping("[TOKEN] [TOKEN]", "##a## ##b##"), {})

    def test_pagination_preserves_styles_assets_conditions_and_currency(self):
        self.config["contentWidth"] = 170
        self.config["bottomY"] = 100
        app.write_json(self.root / "config.json", self.config)
        self.english["rules"]["pages"] = [{"titles": [{"text": "GAME RULES", "color": "gold"}],
            "images": [{"path": "overlay"}], "items": {"setting": {"font": {"size": 21}},
            "lines": ["{fixedJackpot}• A win pays #currency#."]}}]
        self.save_english()
        self.add_german([("A win pays #currency#.", "Ein sehr langer Gewinn wird in der Währung #currency# ausgezahlt.")])
        self.assertEqual(self.run_app(), 0)
        pages = self.generated()["rules"]["pages"]
        self.assertGreater(len(pages), 1)
        for page in pages:
            self.assertEqual(page["items"]["setting"], self.english["rules"]["pages"][0]["items"]["setting"])
            self.assertEqual(page["images"], [{"path": "overlay"}])
            self.assertEqual(page["titles"][0]["color"], "gold")
            self.assertTrue(all(line.startswith("{fixedJackpot}") for line in page["items"]["lines"]))
        self.assertEqual(" ".join(line for page in pages for line in page["items"]["lines"]).count("#currency#"), 1)

    def test_pagination_keeps_blank_lines_only_at_template_positions(self):
        template = game_template()
        template["rules"]["pages"] = [
            {"titles": [{"text": "GAME RULES"}], "items": {"lines": ["• One.", "• Two."]}},
            {"titles": [{"text": "GAME RULES"}], "items": {"lines": ["• Three.", "", "", "• Four."]}}]
        section = app.catalog_sections({"info.json": template})[text_processing.identity("GAME RULES")]
        config = {**self.config, "contentWidth": 1000, "bottomY": -150}
        pages = app.paginate_section(template, section, ["Eins.", "Zwei.", "Drei.", "Vier."],
                                     "REGELN", config, "deu", [])
        self.assertEqual(len(pages), 1)
        self.assertEqual(pages[0]["items"]["lines"], ["• Eins.", "• Zwei.", "• Drei.", "", "", "• Vier."])

    def test_template_layout_moves_midpage_subtitle_with_its_rules(self):
        template = game_template()
        template["rules"]["pages"] = [
            {"titles": [{"text": "FEATURE"}, {"text": "FIRST", "offset": "-500,180", "color": "blue"}],
             "items": {"lines": ["", "", "• One."]}},
            {"titles": [{"text": "FEATURE"}, {"text": "SECOND", "offset": "-500,114", "color": "green"}],
             "items": {"lines": ["• Two.", "", "• Three."]}},
            {"titles": [{"text": "FEATURE"}], "items": {"lines": ["• Four."]}}]
        section = app.catalog_sections({"info.json": template})[text_processing.identity("FEATURE")]
        config = {**self.config, "layout": "template", "contentWidth": 1000}
        pages = app.paginate_section(template, section, ["Eins.", "Zwei.", "Drei.", "Vier."],
                                     "FUNKTION", config, "deu", [], {"first": "ERSTE", "second": "ZWEITE"})
        self.assertEqual(len(pages), 2)
        self.assertEqual(pages[0]["items"]["lines"], ["", "", "• Eins.", "", "• Zwei."])
        self.assertEqual(pages[1]["items"]["lines"], ["", "", "• Drei.", "", "• Vier."])
        self.assertEqual(pages[1]["titles"][1], {"text": "ZWEITE", "offset": "-500,180", "color": "green"})

    def test_template_layout_uses_authored_width_and_capacity(self):
        template = game_template()
        template["rules"]["pages"] = [{"titles": [{"text": "FEATURE"}], "items": {"lines": [
            "• A long authored rule that fits in this game's English layout.", "", "• A final rule."]}}]
        section = app.catalog_sections({"info.json": template})[text_processing.identity("FEATURE")]
        config = {**self.config, "layout": "template", "contentWidth": 50, "bottomY": 170}
        pages = app.paginate_section(template, section,
            ["A translated rule that fits this game's authored space.", "A translated end."],
            "FUNKTION", config, "deu", [])
        self.assertEqual(len(pages), 1)
        self.assertEqual(sum("•" in line for line in pages[0]["items"]["lines"]), 2)

    def test_docx_title_keeps_authored_line_break(self):
        path = self.documents / "GERMAN.docx"
        write_docx(path, [("GAME RULES", "ERSTER TITEL\nZWEITER TITEL", [("First.", "Erste.")])])
        extracted = extraction.extract_docx(path, "deu", [])
        self.assertEqual(extracted[text_processing.identity("GAME RULES")]["title"], "ERSTER TITEL\nZWEITER TITEL")

    def test_blank_separator_does_not_create_empty_page_at_capacity(self):
        template = game_template()
        section = app.catalog_sections({"info.json": template})[text_processing.identity("GAME RULES")]
        config = {**self.config, "contentWidth": 1000, "bottomY": 170}
        pages = app.paginate_section(template, section, ["Eins.", "Zwei.", "Drei."],
                                     "REGELN", config, "deu", [])
        self.assertEqual([page["items"]["lines"] for page in pages], [["• Eins."], ["• Zwei."], ["• Drei."]])

    def test_no_optional_jackpot_template_is_supported(self):
        (self.root / "template/jackpotpages.json").unlink()
        self.add_german()
        self.assertEqual(self.run_app(), 0)
        self.assertFalse((self.root / "output/lang/deu/info/jackpotpages.json").exists())

    def test_missing_english_is_an_actionable_error(self):
        (self.root / "template/info.json").unlink()
        with self.assertRaisesRegex(ValueError, "plantilla inglesa obligatoria"):
            self.run_app()

    def test_missing_documents_report_errors(self):
        self.assertEqual(self.run_app("--lang", "deu"), 1)
        self.assertFalse((self.templates / "deu").exists())
        self.assertFalse((self.root / "output/backups").exists())
        self.assertEqual(self.run_app(), 1)
        self.assertTrue((self.root / "output/templates/info.json").exists())

    def test_generation_preserves_templates_and_apply_is_removed(self):
        self.add_german()
        original = (self.templates / "info.json").read_bytes()
        self.assertEqual(self.run_app(), 0)
        self.assertEqual((self.templates / "info.json").read_bytes(), original)
        self.assertFalse((self.templates / "deu").exists())
        with self.assertRaises(SystemExit) as error:
            self.run_app("--apply")
        self.assertEqual(error.exception.code, 2)

    def test_corrupt_docx_reports_error(self):
        self.add_german()
        (self.documents / "FRENCH.docx").write_text("invalid archive", encoding="utf-8")
        self.assertEqual(self.run_app(), 1)
        self.assertFalse((self.templates / "deu").exists())
        self.assertFalse((self.root / "output/backups").exists())

    def test_reset_preserves_inputs_and_rejects_protected_paths(self):
        self.add_german()
        self.assertEqual(self.run_app(), 0)
        self.assertEqual(self.run_app("--reset"), 0)
        self.assertFalse((self.root / "output").exists())
        self.assertTrue((self.documents / "GERMANI.docx").exists())
        self.assertTrue((self.templates / "info.json").exists())
        self.assertEqual(self.run_app("--reset"), 0)
        with self.assertRaises(ValueError):
            app.reset_output(self.project, self.root, self.documents, self.templates)
        with self.assertRaises(ValueError):
            app.reset_output(self.root / "output", self.root, self.root / "output/documents", self.templates)

    def test_compare_english_false_inserts_every_docx_text_in_order(self):
        self.config["compareEnglish"] = False
        app.write_json(self.root / "config.json", self.config)
        # English column differs from the template and the DOCX has more rows than the template rules.
        self.add_german([("Something else", "• Erstes [SYMBOL_8A] gewinnt."),
                         ("", "Zweite Regel ohne Vergleich."),
                         ("• Final rule.", "Final rule."),
                         ("Extra", "• Zusätzliche Regel.")])
        self.assertEqual(self.run_app(), 0)
        page = self.generated()["rules"]["pages"][0]
        self.assertEqual(page["titles"][0]["text"], "SPIELREGELN")
        rules = [line for line in page["items"]["lines"] if line.strip()]
        self.assertEqual(rules, ["• Erstes ##sym_8a.png## gewinnt.", "• Zweite Regel ohne Vergleich.",
                                 "• Final rule.", "• Zusätzliche Regel."])
        self.assertTrue(any("4 textos" in entry["message"] for entry in self.report()))

    def test_compare_english_false_keeps_unknown_markers_and_layout_capacity(self):
        self.config["compareEnglish"] = False
        app.write_json(self.root / "config.json", self.config)
        self.add_german([("x", f"• Regel {i} [UNBEKANNT].") for i in range(1, 40)])
        self.assertEqual(self.run_app(), 0)
        pages = [page for page in self.generated()["rules"]["pages"]
                 if page.get("items") and page["titles"][0]["text"] == "SPIELREGELN"]
        texts = [line for page in pages for line in page["items"]["lines"] if line.strip()]
        self.assertEqual(len([line for line in texts if line.startswith("•")]), 39)
        self.assertIn("[UNBEKANNT]", texts[0])
        self.assertGreater(len(pages), 2)
        self.assertTrue(any("se inserta tal cual" in entry["message"] for entry in self.report()))


if __name__ == "__main__":
    unittest.main()
